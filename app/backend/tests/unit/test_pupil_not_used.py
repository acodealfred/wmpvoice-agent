"""Pupil dilation is still captured and stored, but must not reach any read-side
consumer: AI report prompts, recovery routing, or the admin exports."""
import csv
import io

from aiohttp import web
from aiohttp.test_utils import TestClient, TestServer
from openpyxl import load_workbook

from ciq.api import history_routes
from ciq.api.history_routes import _PILOT_EXPORT_HEADERS, _TIMELINE_EXPORT_HEADERS
from ciq.recovery.router import (
    _biometric_elevated,
    get_preliminary_recovery_recommendation,
)
from ciq.reports.prompts import (
    RESEARCH_RULES,
    build_analysis_prompt,
    build_biometric_facts,
    build_readiness_analysis_prompt,
    build_snapshot_summary,
)
from survey_loader import load_survey

_SNAPSHOTS = [
    {"questionId": "q1", "domain": "Emotional Exhaustion", "score": 4, "blinkRateChange": 5.0,
     "pupilMmChange": 0.5, "leftGazePosition": "Center", "rightGazePosition": "Center"},
    {"questionId": "q2", "domain": "Depersonalization", "score": 3, "blinkRateChange": -5.0,
     "pupilMmChange": 0.4, "leftGazePosition": "Left", "rightGazePosition": "Center"},
]


def test_report_prompts_never_mention_pupil():
    config = load_survey("TEST")
    for text in (
        RESEARCH_RULES,
        build_biometric_facts(_SNAPSHOTS),
        build_snapshot_summary(_SNAPSHOTS),
        build_analysis_prompt(config, _SNAPSHOTS),
        build_readiness_analysis_prompt(config, _SNAPSHOTS),
    ):
        assert "pupil" not in text.lower()


def test_elevated_pupil_alone_does_not_trigger_recovery_routing():
    assert _biometric_elevated({"blink_band": "Normal", "pupil_band": "High"}) is False
    assert _biometric_elevated({"blink_band": "Elevated (above baseline)"}) is True
    assert _biometric_elevated({"blink_band": "High (below baseline)"}) is True

    result = get_preliminary_recovery_recommendation(
        {"cbi_risk_level": None, "overall_risk_level": None},
        {"blink_band": "Normal", "pupil_band": "High"},
    )
    assert result["track"] != "mindfulness_downshift"


def test_exports_have_no_pupil_columns():
    for headers in (_TIMELINE_EXPORT_HEADERS, _PILOT_EXPORT_HEADERS):
        assert not any("pupil" in h.lower() for h in headers)


_FRAME = {
    "frame_index": 0, "timestamp": "2026-01-01T00:00:00", "turn_state": "user_answer",
    "question_id": "q1", "question_index": 1, "blink_rate_bpm": 15.0,
    "blink_rate_change_percent": 5.0, "pupil_mm_change": 0.42,
    "left_gaze_position": "Center", "right_gaze_position": "Left", "face_emotion": "CALM",
    "voice_sentiment": "neutral", "stress_state": "normal",
}


async def _get(handler, url: str) -> bytes:
    """Call an export handler over a bare app (no auth middleware) and return the body."""
    app = web.Application()
    app.router.add_get(url.split("?")[0], handler)
    async with TestClient(TestServer(app)) as client:
        resp = await client.get(url)
        assert resp.status == 200
        return await resp.read()


async def test_timeline_export_rows_line_up_with_headers(monkeypatch):
    async def frames(_run_id):
        return [_FRAME]

    async def no_record(_run_id):
        return None

    async def no_answers(_run_id):
        return []

    monkeypatch.setattr(history_routes, "get_survey_timeline_frames", frames)
    monkeypatch.setattr(history_routes, "get_survey_record", no_record)
    monkeypatch.setattr(history_routes, "get_survey_item_responses", no_answers)
    body = (await _get(history_routes.admin_export_survey_timeline, "/admin/survey-timeline/export?survey_run_id=run-1")).decode()

    header, row = list(csv.reader(io.StringIO(body)))[:2]
    assert header == _TIMELINE_EXPORT_HEADERS
    assert len(row) == len(header)
    by_name = dict(zip(header, row))
    assert by_name["blink_rate_change_percent"] == "5.0"
    assert by_name["left_gaze_position"] == "Center"
    assert by_name["stress_state"] == "normal"
    assert (by_name["stage"], by_name["segment_index"]) == ("survey", "0")  # legacy row: no stage stored
    assert "0.42" not in row  # stored pupil value never leaks into the export


async def test_pilot_export_rows_line_up_with_headers(monkeypatch):
    async def rows():
        return [{
            "survey_run_id": "run-1", "status": "completed",
            "baseline_blink_rate": 16.5, "baseline_completed_at": "2026-01-01T09:06:00",
            "bat4_total": 12, "bat4_average": 3.0, "cbi3_total": 150, "cbi3_average": 50.0,
            "blink_rate_before": 14.0, "blink_rate_after": 18.0,
            "pupil_dilation_before": 0.11, "pupil_dilation_after": 0.22,
            "response_latency_ms": 2100,
        }]

    monkeypatch.setattr(history_routes, "get_pilot_survey_export_rows", rows)
    body = await _get(history_routes.admin_export_pilot_survey, "/admin/pilot-survey/export")

    ws = load_workbook(io.BytesIO(body)).active
    header, row = [list(r) for r in ws.iter_rows(min_row=1, max_row=2, values_only=True)]
    assert header == _PILOT_EXPORT_HEADERS
    assert len(row) == len(header)
    by_name = dict(zip(header, row))
    assert by_name["Response Latency (ms)"] == 2100
    assert (by_name["Survey Run ID"], by_name["Baseline Blink Rate"]) == ("run-1", 16.5)
    assert 0.11 not in row and 0.22 not in row
